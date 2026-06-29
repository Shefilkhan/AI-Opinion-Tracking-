"""Generate SSRIP Review Report (DOCX) and Presentation (PPTX)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt as PptPt

OUT_DIR = Path(__file__).resolve().parent

PROJECT = {
    "title": "OpinionPulse — AI-Powered Public Opinion Tracking Platform",
    "mentors": "Prof. Kiran Kamlesh Panchal, IOT, Ganpat University",
    "students": "Shefil Khan, Khush Patel",
    "faculty": "Prof. Kiran Kamlesh Panchal",
    "institute": "Conestoga College, Kitchener, Ontario, Canada",
    "start_date": "19th May 2026",
    "review_period": "19th July 2026",
}

# Eight-week narrative reports (GUNI-SSRIP format — paragraph style per week)
WEEKS = [
    {
        "label": "Week 1 Report",
        "report": (
            "In the first week, the main goal was to understand the project scope and define what "
            "OpinionPulse needs to deliver. We studied how public opinion appears across social media "
            "(Reddit, YouTube, Mastodon, Bluesky), news platforms (NewsAPI, Guardian, GNews), and "
            "developer communities (Hacker News, GitHub, Stack Overflow). We reviewed existing sentiment "
            "tracking tools and identified gaps — most tools cover only one platform or lack real-time "
            "AI summaries. We listed core requirements: multi-source keyword search, sentiment scoring, "
            "a unified dashboard, user authentication, and an AI assistant. We drew use-case and data-flow "
            "diagrams showing how a user query travels from the React frontend through FastAPI to external "
            "APIs and back. By the end of this week, we had a clear project plan, technology choices "
            "(React, FastAPI, MySQL), and a shared Git repository structure for the team."
        ),
    },
    {
        "label": "Week 2 Report",
        "report": (
            "The second week focused on setting up the full development environment. We installed and "
            "configured the frontend using React 18, TypeScript, and Vite with Tailwind CSS for styling. "
            "On the backend, we set up Python FastAPI with SQLAlchemy ORM and connected it to a MySQL "
            "database running on XAMPP. We created the initial folder structure — app/api/routes, "
            "app/services, app/schemas, and app/db/models — so frontend and backend code stay organized. "
            "We registered accounts for Reddit, NewsAPI, Guardian, YouTube Data API, and other free "
            "data sources, storing all API keys securely in a .env file (never committed to Git). We "
            "built a health-check endpoint and verified that the frontend proxy correctly forwards "
            "requests to the backend on port 8000. By the end of this week, both servers were running "
            "locally and the team could push code to the Dev branch without conflicts."
        ),
    },
    {
        "label": "Week 3 Report",
        "report": (
            "This week we built the user authentication and account management system. We implemented "
            "user registration with email and password, bcrypt password hashing, and JWT access tokens "
            "for session management. Email OTP verification was added using Gmail SMTP — a 6-digit code "
            "is sent on signup and login to confirm the user's identity. We also built password reset "
            "with secure token links, account lockout after repeated failed login attempts, and a My "
            "Account page where users can update their profile and avatar. On the backend, we created "
            "MySQL tables for users, plans, and search history with SQLAlchemy models and migration "
            "scripts. The frontend received dedicated auth pages: Sign Up, Sign In, Verify OTP, Forgot "
            "Password, and Reset Password. By the end of this week, users could create accounts, verify "
            "their email, and log in securely to access protected dashboard routes."
        ),
    },
    {
        "label": "Week 4 Report",
        "report": (
            "In the fourth week, we built the data collection pipeline that fetches posts and articles "
            "from multiple platforms. We connected the app to the Reddit API, YouTube Data API, and "
            "GDELT news feed so it can retrieve content using any keyword the user enters. Since each "
            "platform returns data in a different format, we wrote a platform_common module to normalize "
            "all results into one structure: title, content, author, platform name, posted date, URL, "
            "and engagement metrics (likes, comments, upvotes). We added deduplication by URL and "
            "headline-focused filtering so only relevant results are shown. Text cleaning removes "
            "excess whitespace and normalizes timestamps to UTC. We stored collected mentions in MySQL "
            "and exposed REST endpoints for projects, keywords, sources, and mentions. By the end of "
            "this week, the backend could fetch and normalize data from the first set of live sources."
        ),
    },
    {
        "label": "Week 5 Report",
        "report": (
            "The fifth week was dedicated to building the sentiment analysis engine and analytics layer. "
            "We integrated VADER (Valence Aware Dictionary and sEntiment Reasoner) to classify each "
            "post as positive, negative, or neutral based on its text content. We added keyword-based "
            "scoring rules for domain-specific terms that VADER might miss. The analytics service "
            "calculates aggregate sentiment percentages, platform breakdowns, and a 24-hour sentiment "
            "trend chart. We built sentiment forecast logic to predict whether opinion is likely to "
            "shift positive or negative over the next seven days. Search history is saved per user so "
            "past queries can be reopened quickly. API endpoints for /api/search and /api/analytics "
            "return structured JSON that the frontend charts with Recharts. All models were tested "
            "with sample queries (e.g. climate change, artificial intelligence) and results were "
            "verified manually for accuracy."
        ),
    },
    {
        "label": "Week 6 Report",
        "report": (
            "This week we built the main frontend that users interact with daily. We created a "
            "BankDash-style dashboard with sidebar navigation, overview cards showing sentiment balance "
            "and trending topic counts, and a weekly activity bar chart. The Search page lets users "
            "enter a keyword, choose platform filters (all, news, tech, Reddit, YouTube), and select "
            "a time range (24 hours, 7 days, 30 days). Results appear as cards with platform badges, "
            "sentiment tags, and direct links to the original source. We added a Wikipedia summary "
            "panel for background context on any topic. The landing page was redesigned with an "
            "editorial style — hero section, feature blocks, tech stack overview, pricing table, and "
            "a public Explore sandbox where visitors can try sample topics without signing up. Settings "
            "and My Account pages were connected to the backend. All pages were tested for responsive "
            "layout and correct data loading without errors."
        ),
    },
    {
        "label": "Week 7 Report",
        "report": (
            "The seventh week focused on advanced features and AI integration. We expanded live search "
            "from a handful of sources to 13 platforms — adding NewsAPI, Guardian, GNews, Currents, "
            "MediaStack, Hacker News, Dev.to, GitHub, Stack Overflow, Mastodon, and Bluesky — all "
            "queried in parallel with a 3-minute cache to respect API rate limits. We built the Compare "
            "Topics module where users enter two keywords (e.g. Tesla vs BYD) and see side-by-side "
            "sentiment charts and word clouds. Pulse AI chat was integrated using Groq (Llama 3.3 70B) — "
            "before answering, the chatbot fetches live search results and uses them as context so "
            "responses are grounded in real data, not generic text. Google OAuth sign-in was added for "
            "one-click authentication. Pro-tier features include Claude-powered risk scoring, debate "
            "analysis, trend prediction, CSV/PDF report export, and personal alert rules. Subscription "
            "plans (Starter and Pro) with server-side usage limits were implemented along with an admin "
            "script to grant Pro access for demo accounts."
        ),
    },
    {
        "label": "Week 8 Report",
        "report": (
            "The last week was used to test, polish, and prepare everything for submission. We ran "
            "end-to-end tests on every feature — registration, login, OAuth, live search across all "
            "13 sources, Compare mode, Pulse AI chat, dashboard widgets, and report export. We fixed "
            "bugs found during testing: Google OAuth invalid_client errors from placeholder env values, "
            "Python 3.9 compatibility issues in Pydantic type hints, news API date-format errors, and "
            "Pulse AI misclassifying sports score queries as opinion comparisons. Team members merged "
            "parallel branches on the Dev branch and resolved conflicts in chat_service.py through "
            "structured code review. We added Terms of Service and Privacy Policy pages, tuned news "
            "APIs for newest-first results, and wrote backend README documentation with setup "
            "instructions. The project code was pushed to GitHub, a professional SSRIP presentation "
            "with UI screenshots was prepared, and this eight-week review report was completed. "
            "Everything was reviewed one final time to ensure all requirements are met before submission."
        ),
    },
]

FUTURE_WORK = (
    "In the future, we plan to deploy OpinionPulse to cloud hosting (AWS or Azure) with production "
    "MySQL, HTTPS, and a CI/CD pipeline. We will add Stripe payment integration for automated Pro and "
    "Enterprise billing. Sentiment accuracy will be improved by replacing keyword-only scoring with "
    "LLM-based analysis for nuanced topics. We plan to build a mobile-responsive PWA and email alert "
    "notifications when sentiment crosses user-defined thresholds. Multi-language support (Hindi, "
    "French, Spanish) will allow analysis of non-English posts. A public REST API with documentation "
    "will enable third-party research teams to integrate OpinionPulse data into their own tools."
)

CHALLENGES = (
    "Managing rate limits across many free news and social APIs required 3-minute caching and parallel "
    "fetch with timeouts. Google OAuth and Gmail SMTP setup needed careful .env configuration for each "
    "teammate's local environment. Merging parallel team branches caused merge conflicts in "
    "chat_service.py — resolved with structured code review. Support needed: stable cloud hosting "
    "credits and Groq/Anthropic API budget for demo presentations and production deployment."
)

COMMENTS = (
    "Students demonstrated strong full-stack progress — from project planning to a working multi-source "
    "AI platform in eight weeks. Recommend demonstrating live search, Compare Topics, and Pulse AI "
    "during the final evaluation. Continue adding automated tests before production deployment."
)

SIGNATURES = [
    ("Name & Sign of Mentor:", "Prof. Kiran Panchal,\nIOT-IT,\nGanpat University"),
    ("Name & Sign of Co-Mentor:", "-"),
    (
        "Name & Sign of Associate Dean (Research)",
        "Prof. (Dr.) Bhavesh Thakar,\nFODE - H&S Department,\nGanpat University",
    ),
    (
        "Name & Sign of Dean / Principal",
        "Prof. (Dr.) K.P. Patel, Principal\nand Executive Dean FODE,\nGanpat University",
    ),
]


def build_docx() -> Path:
    """Build GUNI-SSRIP Review Report matching the official PDF format."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "GUNI Summer Student Research Internship Program\n(GUNI-SSRIP 2026)\nReview Report"
    )
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph()

    fields = [
        ("Project Title:", PROJECT["title"]),
        ("Project Mentor(s):", PROJECT["mentors"]),
        ("Name of Student:", PROJECT["students"]),
        ("Faculty Name:", PROJECT["faculty"]),
        ("Institute Name / Address:", PROJECT["institute"]),
        ("Start Date of SSRP Project:", PROJECT["start_date"]),
        ("Review Period (Date):", PROJECT["review_period"]),
    ]

    for label, value in fields:
        p = doc.add_paragraph()
        p.add_run(f"{label} ").bold = True
        p.add_run(value)

    doc.add_paragraph()
    h = doc.add_paragraph()
    h.add_run(
        "Details of Work done (Weekly) (Please attach detailed report if necessary):"
    ).bold = True

    for week in WEEKS:
        doc.add_paragraph()
        wp = doc.add_paragraph()
        wp.add_run(f"{week['label']}: ").bold = True
        wp.add_run(week["report"])

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Future Work Plan: ").bold = True
    doc.add_paragraph(FUTURE_WORK)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Comments given by Expert / Suggestion:").bold = True
    doc.add_paragraph(COMMENTS)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Challenges / Support Required: ").bold = True
    doc.add_paragraph(CHALLENGES)

    doc.add_paragraph()
    for label, value in SIGNATURES:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        if value and value != "-":
            doc.add_paragraph(value)
        elif value == "-":
            doc.add_paragraph("-")

    path = OUT_DIR / "OpinionPulse_SSRIP_Review_Report.docx"
    doc.save(path)
    return path


SCREENSHOTS_DIR = OUT_DIR / "screenshots"


def _screenshot(name: str) -> Path | None:
    path = SCREENSHOTS_DIR / f"{name}.png"
    return path if path.is_file() else None


def build_pptx() -> Path:
    """Build a project-focused SSRIP presentation with detailed explanations."""

    class C:
        FOREST = RGBColor(0x1B, 0x4D, 0x3E)
        FOREST_DARK = RGBColor(0x12, 0x35, 0x2B)
        CREAM = RGBColor(0xF7, 0xF5, 0xF0)
        WHITE = RGBColor(0xFF, 0xFF, 0xFF)
        TEXT = RGBColor(0x1A, 0x1F, 0x1C)
        MUTED = RGBColor(0x6B, 0x72, 0x6B)
        ACCENT = RGBColor(0xC4, 0x7A, 0x4A)
        LINE = RGBColor(0xD8, 0xD3, 0xCB)

    FONT = "Calibri"
    W = Inches(13.333)
    H = Inches(7.5)
    FOOTER = "OpinionPulse  ·  GUNI SSRIP 2026  ·  Project Presentation"

    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    slide_num = [0]

    def _fill(shape, color: RGBColor) -> None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()

    def _textbox(
        slide,
        left,
        top,
        width,
        height,
        text: str,
        *,
        size: int = 18,
        bold: bool = False,
        color: RGBColor = C.TEXT,
        align=PP_ALIGN.LEFT,
    ):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = align
        p.font.name = FONT
        p.font.size = PptPt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        return box

    def _footer(slide) -> None:
        _textbox(
            slide,
            Inches(0.6),
            H - Inches(0.45),
            W - Inches(1.2),
            Inches(0.3),
            f"{FOOTER}  ·  Slide {slide_num[0]}",
            size=9,
            color=C.MUTED,
            align=PP_ALIGN.RIGHT,
        )

    def _slide_bg(slide, color: RGBColor = C.CREAM) -> None:
        bg = slide.shapes.add_shape(1, 0, 0, W, H)
        _fill(bg, color)
        slide.shapes._spTree.remove(bg._element)
        slide.shapes._spTree.insert(2, bg._element)

    def _header_bar(slide, title: str, subtitle: str = "") -> None:
        bar = slide.shapes.add_shape(1, 0, 0, W, Inches(1.05))
        _fill(bar, C.FOREST)
        accent = slide.shapes.add_shape(1, 0, Inches(1.05), W, Inches(0.06))
        _fill(accent, C.ACCENT)
        _textbox(slide, Inches(0.65), Inches(0.22), Inches(11.5), Inches(0.55), title, size=28, bold=True, color=C.WHITE)
        if subtitle:
            _textbox(slide, Inches(0.65), Inches(0.72), Inches(11.5), Inches(0.35), subtitle, size=13, color=RGBColor(0xC8, 0xD9, 0xD2))

    def _left_accent(slide) -> None:
        stripe = slide.shapes.add_shape(1, 0, Inches(1.11), Inches(0.12), H - Inches(1.11))
        _fill(stripe, C.FOREST)

    def _bullets(
        slide,
        items: list[str],
        *,
        top=Inches(1.45),
        left=Inches(0.85),
        width=None,
        size: int = 15,
        gap: float = 0.52,
    ) -> None:
        text_width = width or (W - left - Inches(0.9))
        y = top
        for item in items:
            dot = slide.shapes.add_shape(1, left, y + Inches(0.07), Inches(0.1), Inches(0.1))
            _fill(dot, C.ACCENT)
            _textbox(slide, left + Inches(0.24), y, text_width, Inches(0.85), item, size=size, color=C.TEXT)
            y += Inches(gap)

    def _paragraph_block(slide, text: str, *, top=Inches(1.45), left=Inches(0.85), width=Inches(11.8), size: int = 15) -> None:
        _textbox(slide, left, top, width, Inches(1.2), text, size=size, color=C.TEXT)

    def _new_content(title: str, subtitle: str = "") -> object:
        slide_num[0] += 1
        slide = prs.slides.add_slide(blank)
        _slide_bg(slide)
        _header_bar(slide, title, subtitle)
        _left_accent(slide)
        _footer(slide)
        return slide

    def _new_section(title: str, subtitle: str) -> None:
        slide_num[0] += 1
        slide = prs.slides.add_slide(blank)
        _slide_bg(slide, C.FOREST_DARK)
        bar = slide.shapes.add_shape(1, Inches(0.8), Inches(2.8), Inches(1.2), Inches(0.08))
        _fill(bar, C.ACCENT)
        _textbox(slide, Inches(0.8), Inches(3.0), Inches(11), Inches(0.9), title, size=36, bold=True, color=C.WHITE)
        _textbox(slide, Inches(0.8), Inches(3.85), Inches(11), Inches(0.6), subtitle, size=20, color=RGBColor(0xA8, 0xC4, 0xB8))
        _footer(slide)

    def _add_image(slide, img_path: Path | None, left, top, width) -> None:
        frame = slide.shapes.add_shape(1, left - Inches(0.08), top - Inches(0.08), width + Inches(0.16), width * 0.56 + Inches(0.16))
        _fill(frame, C.WHITE)
        frame.line.color.rgb = C.LINE
        if img_path:
            slide.shapes.add_picture(str(img_path), left, top, width=width)
        else:
            _textbox(slide, left, top + Inches(1.0), width, Inches(0.5), "[Screenshot placeholder]", size=12, color=C.MUTED, align=PP_ALIGN.CENTER)

    def _screenshot_slide(title: str, subtitle: str, img_name: str, points: list[str], caption: str = "") -> None:
        slide = _new_content(title, subtitle)
        _bullets(slide, points, top=Inches(1.35), left=Inches(0.75), width=Inches(5.9), size=13, gap=0.46)
        img = _screenshot(img_name)
        _add_image(slide, img, Inches(6.85), Inches(1.3), Inches(6.1))
        if caption:
            _textbox(slide, Inches(6.85), Inches(6.55), Inches(6.1), Inches(0.35), caption, size=10, color=C.MUTED, align=PP_ALIGN.CENTER)

    # ── Title ──
    slide_num[0] += 1
    slide = prs.slides.add_slide(blank)
    panel = slide.shapes.add_shape(1, 0, 0, Inches(5.2), H)
    _fill(panel, C.FOREST_DARK)
    right = slide.shapes.add_shape(1, Inches(5.2), 0, W - Inches(5.2), H)
    _fill(right, C.CREAM)
    line = slide.shapes.add_shape(1, Inches(5.15), Inches(1.8), Inches(0.06), Inches(3.8))
    _fill(line, C.ACCENT)
    _textbox(slide, Inches(0.7), Inches(2.0), Inches(4.2), Inches(0.4), "GUNI SSRIP 2026", size=13, color=RGBColor(0xA8, 0xC4, 0xB8))
    _textbox(slide, Inches(0.7), Inches(2.45), Inches(4.2), Inches(1.2), "OpinionPulse", size=44, bold=True, color=C.WHITE)
    _textbox(slide, Inches(0.7), Inches(3.55), Inches(4.2), Inches(1.0), "AI-Powered Public\nOpinion Tracking Platform", size=18, color=RGBColor(0xD4, 0xE4, 0xDE))
    _textbox(slide, Inches(5.55), Inches(2.35), Inches(7.2), Inches(0.5), "Project Presentation", size=14, bold=True, color=C.MUTED)
    _textbox(
        slide,
        Inches(5.55),
        Inches(2.95),
        Inches(7.2),
        Inches(1.6),
        "Full-stack web platform that aggregates,\nanalyzes, and visualizes public opinion\nfrom 13+ live data sources",
        size=24,
        bold=True,
        color=C.TEXT,
    )
    _textbox(slide, Inches(5.55), Inches(4.75), Inches(7.2), Inches(0.4), PROJECT["institute"], size=14, color=C.MUTED)
    _textbox(slide, Inches(5.55), Inches(5.2), Inches(7.2), Inches(0.5), "Shefil Khan  ·  Khush Patel", size=14, color=C.TEXT)
    _footer(slide)

    # ── Agenda ──
    slide = _new_content("Presentation Agenda", "Project documentation — not weekly progress report")
    _bullets(
        slide,
        [
            "Project introduction, objectives, and scope",
            "System architecture and technology stack",
            "Application modules with UI screenshots",
            "Multi-source data integration and AI pipeline",
            "Database design, security, and subscription model",
            "Future enhancements and conclusion",
        ],
        top=Inches(1.4),
        size=16,
        gap=0.58,
    )

    # ── Project introduction ──
    slide = _new_content("Project Introduction", PROJECT["title"])
    _paragraph_block(
        slide,
        "OpinionPulse is a full-stack web application built during the GUNI Summer Student Research "
        "Internship Program (SSRIP). It helps researchers, journalists, and analysts understand how "
        "the public feels about any topic by searching news articles, social posts, and developer "
        "communities in one place — then scoring sentiment and generating AI summaries.",
        top=Inches(1.35),
    )
    _bullets(
        slide,
        [
            "Purpose: unify scattered online conversations into actionable sentiment intelligence",
            "Users enter a keyword (e.g. climate change, AI regulation) and receive live results within seconds",
            "Each result is tagged positive, neutral, or negative with source links and engagement metrics",
            "Pro users unlock Pulse AI chat, compare mode, risk scoring, and advanced analytics",
        ],
        top=Inches(2.55),
        size=14,
        gap=0.5,
    )

    # ── Problem & objectives ──
    slide = _new_content("Problem Statement & Objectives", "Why this project exists")
    _bullets(
        slide,
        [
            "Problem: public opinion is fragmented across Reddit, YouTube, news sites, GitHub, Hacker News, and more — manual tracking is slow and inconsistent",
            "Objective 1: build a single dashboard that searches all major opinion sources in parallel",
            "Objective 2: apply automated sentiment analysis so users see positive/neutral/negative breakdown instantly",
            "Objective 3: integrate AI (Groq Llama 3.3) to answer natural-language questions using freshly fetched data",
            "Objective 4: provide exportable reports (CSV/PDF) and alerts for ongoing topic monitoring",
        ],
        top=Inches(1.35),
        size=14,
        gap=0.48,
    )

    # ── Scope & users ──
    slide = _new_content("Project Scope & Target Users", "What the platform delivers")
    cols = [
        ("In Scope", [
            "Multi-source keyword search with time and platform filters",
            "Side-by-side topic comparison with sentiment charts",
            "Live dashboard: trending topics, debates, platform pulse",
            "Pulse AI assistant with real search context",
            "User accounts with JWT auth, OTP email, Google OAuth",
            "Starter / Pro subscription tiers with usage limits",
        ]),
        ("Target Users", [
            "Journalists tracking breaking story sentiment",
            "Researchers studying public opinion trends",
            "Marketing teams monitoring brand perception",
            "Policy analysts comparing debate narratives",
            "Students learning full-stack + AI integration",
        ]),
    ]
    for i, (heading, items) in enumerate(cols):
        left = Inches(0.75 + i * 6.35)
        box = slide.shapes.add_shape(1, left, Inches(1.45), Inches(6.0), Inches(5.5))
        _fill(box, C.WHITE)
        box.line.color.rgb = C.LINE
        head = slide.shapes.add_shape(1, left, Inches(1.45), Inches(6.0), Inches(0.5))
        _fill(head, C.FOREST if i == 0 else C.ACCENT)
        _textbox(slide, left + Inches(0.2), Inches(1.52), Inches(5.6), Inches(0.35), heading, size=15, bold=True, color=C.WHITE, align=PP_ALIGN.CENTER)
        _bullets(slide, items, top=Inches(2.15), left=left + Inches(0.25), width=Inches(5.5), size=13, gap=0.44)

    # ── Architecture ──
    slide = _new_content("System Architecture", "Three-tier design with async API aggregation")
    layers = [
        ("Presentation Layer", "React 18 + TypeScript SPA\nPages: Landing, Dashboard, Search, Compare,\nChat, Reports, Alerts, Settings, Explore sandbox"),
        ("Application Layer", "FastAPI REST API\nServices: search, sentiment, chat, dashboard,\nreports, billing, OAuth, email OTP"),
        ("Data Layer", "MySQL (XAMPP) via SQLAlchemy\nTables: users, plans, search_history,\nmentions, chat_sessions, reports, alerts"),
        ("External APIs", "13 platforms queried in parallel\nNewsAPI, Guardian, GNews, Currents,\nMediaStack, Reddit, YouTube, GitHub, etc."),
    ]
    y = Inches(1.35)
    for heading, body in layers:
        card = slide.shapes.add_shape(1, Inches(0.75), y, Inches(11.85), Inches(1.15))
        _fill(card, C.WHITE)
        card.line.color.rgb = C.LINE
        tag = slide.shapes.add_shape(1, Inches(0.75), y, Inches(2.4), Inches(1.15))
        _fill(tag, C.FOREST)
        _textbox(slide, Inches(0.9), y + Inches(0.35), Inches(2.1), Inches(0.5), heading, size=13, bold=True, color=C.WHITE)
        _textbox(slide, Inches(3.35), y + Inches(0.18), Inches(8.9), Inches(0.85), body, size=12, color=C.TEXT)
        y += Inches(1.28)

    # ── Tech stack ──
    slide = _new_content("Technology Stack", "Tools and frameworks used in development")
    stacks = [
        ("Frontend", [
            "React 18 + TypeScript — component-based UI",
            "Vite — fast dev server and production bundler",
            "Tailwind CSS — responsive editorial design system",
            "React Query — cached API state management",
            "Recharts — sentiment charts and trend graphs",
        ]),
        ("Backend", [
            "Python FastAPI — async REST endpoints",
            "SQLAlchemy ORM — MySQL database models",
            "Pydantic — request/response validation",
            "JWT + bcrypt — secure authentication",
            "Gmail SMTP — OTP verification emails",
        ]),
        ("AI & Integrations", [
            "Groq (Llama 3.3 70B) — Pulse AI chat responses",
            "Anthropic Claude — Pro-tier insight generation",
            "VADER + keyword scoring — sentiment engine",
            "13 external APIs — parallel async fetch with 3-min cache",
        ]),
    ]
    for i, (heading, items) in enumerate(stacks):
        left = Inches(0.75 + i * 4.15)
        box = slide.shapes.add_shape(1, left, Inches(1.4), Inches(3.85), Inches(5.6))
        _fill(box, C.WHITE)
        box.line.color.rgb = C.LINE
        head_bar = slide.shapes.add_shape(1, left, Inches(1.4), Inches(3.85), Inches(0.5))
        _fill(head_bar, C.FOREST if i == 0 else (C.ACCENT if i == 1 else RGBColor(0x2D, 0x6A, 0x4E)))
        _textbox(slide, left + Inches(0.2), Inches(1.47), Inches(3.5), Inches(0.35), heading, size=14, bold=True, color=C.WHITE, align=PP_ALIGN.CENTER)
        _bullets(slide, items, top=Inches(2.05), left=left + Inches(0.2), width=Inches(3.45), size=12, gap=0.42)

    _new_section("Application Modules", "Detailed walkthrough with screenshots")

    _screenshot_slide(
        "Landing Page & Onboarding",
        "Public-facing entry point for new users",
        "01_landing_hero",
        [
            "Editorial-style landing page introduces OpinionPulse value proposition",
            "Hero section explains real-time multi-source opinion tracking",
            "Feature blocks describe Search, Compare, Dashboard, and Pulse AI modules",
            "Pricing section shows Starter (free) and Pro plans with feature comparison",
            "Call-to-action routes users to sign-up or the interactive Explore sandbox",
        ],
        caption="Screenshot: Landing page hero section",
    )

    _screenshot_slide(
        "Landing Page — Feature Overview",
        "How the product is explained to visitors",
        "02_landing_features",
        [
            "Three-step workflow: Search any topic → Analyze sentiment → Act on insights",
            "Tech stack section highlights React, FastAPI, Groq AI, and 13 data sources",
            "Social proof and editorial quote reinforce research-grade credibility",
            "Footer links to Terms of Service and Privacy Policy for compliance",
        ],
        caption="Screenshot: Landing page features section",
    )

    _screenshot_slide(
        "Interactive Explore Sandbox",
        "Public demo — no login required",
        "03_explore_search_results",
        [
            "Explore page lets visitors try OpinionPulse before creating an account",
            "Pre-built topics (Bitcoin, AI, Climate) show sample multi-platform results",
            "Each card displays platform badge, author, timestamp, sentiment tag, and engagement",
            "AI summary paragraph explains overall sentiment direction for the topic",
            "Users can also type a custom keyword to run a live API search",
        ],
        caption="Screenshot: Explore sandbox — Bitcoin topic results",
    )

    _screenshot_slide(
        "Dashboard Module",
        "Central hub for sentiment overview and trending activity",
        "04_explore_dashboard",
        [
            "Dashboard shows sentiment balance (% positive), trending topic count, and live source count",
            "Weekly activity bar chart tracks positive vs negative sentiment over 7 days",
            "Trending topics table lists hot keywords with sentiment direction arrows",
            "Platform pulse widget shows which sources (Reddit, YouTube, news) are most active",
            "Debate tracker highlights polarized discussions across platforms",
        ],
        caption="Screenshot: Dashboard preview (Explore sandbox)",
    )

    slide = _new_content("Multi-Source Search Engine", "Core module — queries 13 platforms in parallel")
    _paragraph_block(
        slide,
        "When a user searches a keyword, the FastAPI backend fires async requests to every configured "
        "platform simultaneously. Results are deduplicated, sorted by posted date (newest first), "
        "scored for sentiment, and returned with a Wikipedia summary and 24-hour trend chart.",
        top=Inches(1.3),
        size=13,
    )
    sources_left = [
        "Social: Reddit, YouTube, Mastodon, Bluesky",
        "News: NewsAPI, Guardian, GNews, Currents, MediaStack",
        "Tech: Hacker News, Dev.to, GitHub, Stack Overflow",
        "Reference: Wikipedia summary (context panel)",
    ]
    sources_right = [
        "Filters: platform type, time range (24h / 7d / 30d / all)",
        "Headline-focused matching — results must contain query keywords",
        "3-minute in-memory cache reduces API rate-limit hits",
        "Search history saved per user for quick re-access",
        "Sentiment forecast predicts 7-day trend direction",
    ]
    _bullets(slide, sources_left, top=Inches(2.35), left=Inches(0.75), width=Inches(5.8), size=12, gap=0.42)
    _bullets(slide, sources_right, top=Inches(2.35), left=Inches(6.85), width=Inches(5.8), size=12, gap=0.42)
    img = _screenshot("05_explore_ai_topic")
    _add_image(slide, img, Inches(6.85), Inches(5.0), Inches(5.8))

    slide = _new_content("Compare Topics Module", "Side-by-side sentiment analysis for two keywords")
    _bullets(
        slide,
        [
            "User enters Topic A and Topic B (e.g. Tesla vs BYD, Bitcoin vs Ethereum)",
            "Backend runs parallel searches for both topics across all enabled platforms",
            "Results page shows dual sentiment pie charts, platform breakdown, and keyword word clouds",
            "Comparison table highlights which topic has higher positive/negative ratio per platform",
            "Useful for brand rivalry analysis, election tracking, and product launch monitoring",
            "Pro plan required for unlimited compare searches per month",
        ],
        top=Inches(1.35),
        size=14,
        gap=0.46,
    )

    slide = _new_content("Pulse AI Assistant", "Groq-powered conversational intelligence")
    _bullets(
        slide,
        [
            "Floating chat bubble available on every authenticated page",
            "Dedicated /chat page for full-screen conversation history",
            "Before answering, Pulse AI fetches live search results for the user's question",
            "Question classifier routes queries: factual, compare, sentiment, trend, or general",
            "Sports/score queries are detected separately to avoid false compare-mode triggers",
            "Groq Llama 3.3 70B generates concise, data-grounded responses (not generic chatbot text)",
            "Pro users get Claude-powered deep insights: risk profiles, debate analysis, crisis response",
        ],
        top=Inches(1.35),
        size=13,
        gap=0.44,
    )

    slide = _new_content("Analytics, Reports & Alerts", "Turning raw data into actionable output")
    modules = [
        ("Analytics", "Word cloud from extracted keywords, platform sentiment bar chart, 24h trend line graph, sentiment forecast for next 7 days"),
        ("Reports", "Generate CSV or PDF exports of search results with sentiment scores, source URLs, and timestamps. Search history archive for re-download."),
        ("Alerts", "Configure keyword alert rules with sentiment thresholds. System evaluates new mentions and flags spikes in negative sentiment."),
    ]
    for i, (heading, body) in enumerate(modules):
        top = Inches(1.45 + i * 1.75)
        card = slide.shapes.add_shape(1, Inches(0.75), top, Inches(11.85), Inches(1.5))
        _fill(card, C.WHITE)
        card.line.color.rgb = C.LINE
        _textbox(slide, Inches(0.95), top + Inches(0.15), Inches(2.5), Inches(0.35), heading, size=15, bold=True, color=C.FOREST)
        _textbox(slide, Inches(3.55), top + Inches(0.15), Inches(8.8), Inches(1.1), body, size=13, color=C.TEXT)

    _screenshot_slide(
        "Subscription Plans",
        "Starter vs Pro tier feature gating",
        "06_pricing_plans",
        [
            "Starter (Free): limited searches, basic dashboard, explore sandbox access",
            "Pro ($19/mo): unlimited searches, Compare mode, Pulse AI, PDF export, risk scoring",
            "Enterprise: custom API access and team seats (planned)",
            "Usage limits enforced server-side via plan_id on user record",
            "Admin script (set_user_plan.py) allows manual Pro access for demo accounts",
        ],
        caption="Screenshot: Pricing page with plan comparison",
    )

    slide = _new_content("Authentication & Security", "How user accounts are protected")
    _bullets(
        slide,
        [
            "Registration: email + password with bcrypt hashing and email OTP verification (6-digit code via Gmail SMTP)",
            "Login: JWT access tokens with configurable expiry; failed attempts trigger account lockout",
            "Google OAuth 2.0: one-click sign-in with server-side token validation (placeholder client IDs rejected)",
            "Password reset: secure token link sent to registered email address",
            "All project-scoped API routes require Authorization: Bearer header and verify user ownership",
            "Environment variables (.env) store API keys — never committed to Git repository",
        ],
        top=Inches(1.35),
        size=13,
        gap=0.44,
    )

    slide = _new_content("Database Design", "MySQL schema — core entities")
    entities = [
        ("users", "email, password_hash, google_id, plan_id, role, OTP fields, lockout counters"),
        ("search_history", "user_id, query, platform_filter, result_count, created_at"),
        ("mentions", "project_id, source, content, sentiment_score, posted_at, url"),
        ("chat_sessions / chat_messages", "user_id, session title, role (user/assistant), message content"),
        ("plans", "id (starter/pro), search_limit, compare_limit, features JSON"),
        ("reports / alerts", "user_id, query, format, threshold rules, last_evaluated_at"),
    ]
    y = Inches(1.35)
    for table, cols in entities:
        row = slide.shapes.add_shape(1, Inches(0.75), y, Inches(11.85), Inches(0.82))
        _fill(row, C.WHITE)
        row.line.color.rgb = C.LINE
        _textbox(slide, Inches(0.95), y + Inches(0.18), Inches(2.2), Inches(0.4), table, size=13, bold=True, color=C.FOREST)
        _textbox(slide, Inches(3.2), y + Inches(0.18), Inches(9.2), Inches(0.5), cols, size=12, color=C.TEXT)
        y += Inches(0.92)

    slide = _new_content("Data Flow Pipeline", "From user query to AI response")
    steps = [
        ("1. User Input", "Keyword + filters submitted from React SearchPage via POST /api/search"),
        ("2. Parallel Fetch", "search_service.py dispatches asyncio tasks to 13 platform modules simultaneously"),
        ("3. Normalize", "platform_common.py deduplicates by URL, normalizes fields (title, author, posted_at, engagement)"),
        ("4. Sentiment", "VADER + keyword rules score each result; aggregate summary calculated"),
        ("5. Store & Return", "Results cached 3 min, saved to search_history, returned as JSON to frontend"),
        ("6. AI Context", "chat_service.py reuses search results as context for Groq/Claude prompt"),
    ]
    for i, (step, detail) in enumerate(steps):
        col = i % 2
        row = i // 2
        left = Inches(0.75 + col * 6.2)
        top = Inches(1.4 + row * 1.85)
        card = slide.shapes.add_shape(1, left, top, Inches(5.85), Inches(1.55))
        _fill(card, C.WHITE)
        card.line.color.rgb = C.LINE
        _textbox(slide, left + Inches(0.15), top + Inches(0.12), Inches(5.5), Inches(0.35), step, size=13, bold=True, color=C.FOREST)
        _textbox(slide, left + Inches(0.15), top + Inches(0.5), Inches(5.5), Inches(0.9), detail, size=11, color=C.MUTED)

    slide = _new_content("Future Enhancements", "Planned improvements beyond SSRIP")
    _bullets(
        slide,
        [
            "Cloud deployment (AWS/Azure) with production MySQL, HTTPS, and CI/CD pipeline",
            "Stripe payment integration for automated Pro/Enterprise billing",
            "LLM-based sentiment (replacing keyword/VADER) for nuanced technical topics",
            "Email/push alert notifications when sentiment crosses configured thresholds",
            "Mobile-responsive PWA and native app for on-the-go monitoring",
            "Public REST API with documentation for third-party research integrations",
        ],
        top=Inches(1.4),
        size=14,
        gap=0.5,
    )

    slide = _new_content("Project Outcomes", "What we delivered")
    _bullets(
        slide,
        [
            "Fully functional full-stack web application with 13 live data source integrations",
            "Secure authentication system (JWT, OTP, Google OAuth) with role-based plan limits",
            "AI-powered Pulse chat assistant grounded in real-time search data",
            "Professional editorial UI with landing page, explore sandbox, and BankDash-style dashboard",
            "Complete backend API documented at /docs with health checks and smoke tests",
            "Open-source codebase on GitHub — ready for demo and further development",
        ],
        top=Inches(1.4),
        size=14,
        gap=0.5,
    )

    # ── Closing ──
    slide_num[0] += 1
    slide = prs.slides.add_slide(blank)
    _slide_bg(slide, C.FOREST_DARK)
    _textbox(slide, Inches(0.8), Inches(2.2), Inches(11.5), Inches(0.9), "Thank You", size=44, bold=True, color=C.WHITE, align=PP_ALIGN.CENTER)
    _textbox(slide, Inches(0.8), Inches(3.2), Inches(11.5), Inches(0.6), "OpinionPulse — Real-time public opinion intelligence", size=20, color=RGBColor(0xA8, 0xC4, 0xB8), align=PP_ALIGN.CENTER)
    _textbox(slide, Inches(0.8), Inches(4.0), Inches(11.5), Inches(0.5), "Live demo available  ·  Questions welcome", size=16, color=C.ACCENT, align=PP_ALIGN.CENTER)
    _textbox(slide, Inches(0.8), Inches(4.8), Inches(11.5), Inches(0.4), "github.com/Shefilkhan/AI-Opinion-Tracking-", size=12, color=C.MUTED, align=PP_ALIGN.CENTER)
    _footer(slide)

    path = OUT_DIR / "OpinionPulse_SSRIP_Presentation.pptx"
    prs.save(path)
    return path


if __name__ == "__main__":
    import subprocess
    import sys

    missing = {"01_landing_hero", "02_landing_features", "03_explore_search_results",
               "04_explore_dashboard", "05_explore_ai_topic", "06_pricing_plans"} - {
        p.stem for p in SCREENSHOTS_DIR.glob("*.png")
    }
    if missing:
        print(f"Missing screenshots ({len(missing)}) — capturing from local dev server...")
        cap = Path(__file__).resolve().parent / "capture_screenshots.py"
        result = subprocess.run([sys.executable, str(cap)], check=False)
        if result.returncode != 0:
            print("Warning: some screenshots may be missing; PPT will use placeholders where needed.")

    docx_path = build_docx()
    pptx_path = build_pptx()
    print(f"Created: {docx_path}")
    print(f"Created: {pptx_path}")
